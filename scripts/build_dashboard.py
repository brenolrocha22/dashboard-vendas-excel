import sys
from collections import OrderedDict
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import DataBarRule, CellIsRule
from openpyxl.chart import BarChart, LineChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.series import DataPoint
from openpyxl.drawing.line import LineProperties

SRC, OUT = sys.argv[1], sys.argv[2]

PRETO, AZUL, AZUL_ESC, AZUL_CLARO, CINZA_TXT = "111111", "1565C0", "0D47A1", "E3F2FD", "757575"
PALETA = ["0D47A1", "1565C0", "1E88E5", "42A5F5", "90CAF9", "263238", "607D8B"]
fino = Side(style="thin", color="D0D0D0")
BORDA = Border(left=fino, right=fino, top=fino, bottom=fino)
MOEDA = '"R$" #,##0'
MOEDA_K = '[>=1000000]"R$" #,##0.0,,"M";[>=1000]"R$" #,##0,"mil";"R$" #,##0'
PAISES_PT = {"Brazil": "Brasil", "France": "França", "China": "China", "United States": "Estados Unidos",
             "Japan": "Japão", "Argentina": "Argentina"}
MESES = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]

# ---------- dados ----------
src = openpyxl.load_workbook(SRC, read_only=True)["base"]
linhas = [r for r in src.iter_rows(min_row=2, values_only=True) if r[0] is not None]
# alguns modelos vêm como número (ex.: 500); padroniza tudo como texto
linhas = [r[:2] + (str(r[2]).strip(),) + r[3:] for r in linhas]
linhas.sort(key=lambda r: (r[7], r[0]))
marcas = sorted({r[1] for r in linhas})
paises = sorted({PAISES_PT.get(r[4], r[4]) for r in linhas})
anos = sorted({r[7].year for r in linhas})
modelo_marca = OrderedDict(sorted({(r[2], r[1]) for r in linhas}))
N = len(linhas)
ULT = N + 1

wb = Workbook()
dash = wb.active
dash.title = "Dashboard"
base = wb.create_sheet("Base")
calc = wb.create_sheet("Cálculos")


def nome(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


# ---------- Base (tabela) ----------
cab = ["ID", "Data", "Ano", "Mês", "Marca", "Modelo", "Preço", "País", "Região", "Cliente",
       "Filtro atual", "Filtro ano anterior", "Filtro sem ano"]
base.append(cab)
for i, r in enumerate(linhas):
    lin = i + 2
    base.append([r[0], r[7], f"=YEAR(B{lin})", f"=MONTH(B{lin})", r[1], r[2], r[3], PAISES_PT.get(r[4], r[4]), r[5], r[6],
                 f'=IF(AND(OR(sel_ano="Todos",C{lin}=sel_ano),OR(sel_pais="Todos",H{lin}=sel_pais),OR(sel_marca="Todas",E{lin}=sel_marca)),1,0)',
                 f'=IF(sel_ano="Todos",0,IF(AND(C{lin}=sel_ano-1,OR(sel_pais="Todos",H{lin}=sel_pais),OR(sel_marca="Todas",E{lin}=sel_marca)),1,0))',
                 f'=IF(AND(OR(sel_pais="Todos",H{lin}=sel_pais),OR(sel_marca="Todas",E{lin}=sel_marca)),1,0)'])
for lin in range(2, ULT + 1):
    base[f"B{lin}"].number_format = "dd/mm/yyyy"
    base[f"G{lin}"].number_format = '"R$" #,##0.00'
tab = Table(displayName="tbVendas", ref=f"A1:M{ULT}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
base.add_table(tab)
for col, w in zip("ABCDEFGHIJKLM", [7, 12, 7, 6, 13, 16, 14, 16, 18, 24, 10, 12, 10]):
    base.column_dimensions[col].width = w
base.freeze_panes = "A2"
rg = lambda c: f"Base!${c}$2:${c}${ULT}"
PRECO, ANO, MES, MARCA, MODELO, PAIS, FA, FANT, FSA = (rg(c) for c in "GCDEFHKLM")

# ---------- Cálculos ----------
calc.sheet_view.showGridLines = False
calc["A1"].value = "Aba de apoio: tudo aqui é calculado a partir da Base e dos filtros do Dashboard."
calc["A1"].font = Font(italic=True, color=CINZA_TXT)


def bloco(col, linha, titulo, cabecalhos):
    calc.cell(row=linha, column=col, value=titulo).font = Font(bold=True, color=AZUL)
    for k, h in enumerate(cabecalhos):
        c = calc.cell(row=linha + 1, column=col + k, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
    return linha + 2


def L(col):
    return openpyxl.utils.get_column_letter(col)


# listas dos filtros (col A..C)
r0 = bloco(1, 3, "Listas dos filtros", ["Anos", "Países", "Marcas"])
for k, v in enumerate(["Todos"] + anos):
    calc.cell(row=r0 + k, column=1, value=v)
for k, v in enumerate(["Todos"] + paises):
    calc.cell(row=r0 + k, column=2, value=v)
for k, v in enumerate(["Todas"] + marcas):
    calc.cell(row=r0 + k, column=3, value=v)
nome("lista_anos", f"'Cálculos'!$A${r0}:$A${r0 + len(anos)}")
nome("lista_paises", f"'Cálculos'!$B${r0}:$B${r0 + len(paises)}")
nome("lista_marcas", f"'Cálculos'!$C${r0}:$C${r0 + len(marcas)}")

# KPIs (col E..F)
r = bloco(5, 3, "KPIs", ["Indicador", "Valor"])
kpis = [
    ("Faturamento", f"=SUMIFS({PRECO},{FA},1)", "kpi_fat"),
    ("Vendas", f"=COUNTIFS({FA},1)", "kpi_qtd"),
    ("Ticket médio", "=IFERROR(kpi_fat/kpi_qtd,0)", "kpi_ticket"),
    ("Faturamento ano anterior", f"=SUMIFS({PRECO},{FANT},1)", "kpi_fat_ant"),
    ("Variação vs ano anterior", '=IF(OR(sel_ano="Todos",kpi_fat_ant=0),"",kpi_fat/kpi_fat_ant-1)', "kpi_var"),
]
for k, (rot, f, nm) in enumerate(kpis):
    calc.cell(row=r + k, column=5, value=rot)
    calc.cell(row=r + k, column=6, value=f)
    nome(nm, f"'Cálculos'!$F${r + k}")


def ranking(col, linha, titulo, itens, criterio_col, extra=None):
    """Tabela item | faturamento | qtd | ajuste (desempate) e versão ordenada ao lado."""
    r = bloco(col, linha, titulo, ["Item", "Faturamento", "Vendas", "Ajuste"] + (["Marca"] if extra else []))
    n = len(itens)
    for k, it in enumerate(itens):
        lin = r + k
        calc.cell(row=lin, column=col, value=it)
        calc.cell(row=lin, column=col + 1, value=f"=SUMIFS({PRECO},{criterio_col},{L(col)}{lin},{FA},1)")
        calc.cell(row=lin, column=col + 2, value=f"=COUNTIFS({criterio_col},{L(col)}{lin},{FA},1)")
        calc.cell(row=lin, column=col + 3, value=f"={L(col + 1)}{lin}+ROW()/1000000")
        if extra:
            calc.cell(row=lin, column=col + 4, value=extra[it])
    itens_rg = f"${L(col)}${r}:${L(col)}${r + n - 1}"
    adj_rg = f"${L(col + 3)}${r}:${L(col + 3)}${r + n - 1}"
    fat_rg = f"${L(col + 1)}${r}:${L(col + 1)}${r + n - 1}"
    qtd_rg = f"${L(col + 2)}${r}:${L(col + 2)}${r + n - 1}"
    oc = col + (6 if extra else 5)
    calc.cell(row=linha, column=oc, value=titulo + " (ordenado)").font = Font(bold=True, color=AZUL)
    for k, h in enumerate(["#", "Item", "Faturamento", "Vendas", "% do total"] + (["Marca"] if extra else [])):
        c = calc.cell(row=linha + 1, column=oc + k, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
    for k in range(n):
        lin = r + k
        calc.cell(row=lin, column=oc, value=k + 1)
        pos = f"MATCH(LARGE({adj_rg},{L(oc)}{lin}),{adj_rg},0)"
        calc.cell(row=lin, column=oc + 1, value=f"=INDEX({itens_rg},{pos})")
        calc.cell(row=lin, column=oc + 2, value=f"=INDEX({fat_rg},{pos})").number_format = MOEDA
        calc.cell(row=lin, column=oc + 3, value=f"=INDEX({qtd_rg},{pos})")
        calc.cell(row=lin, column=oc + 4, value=f"=IFERROR({L(oc + 2)}{lin}/kpi_fat,0)").number_format = "0.0%"
        if extra:
            calc.cell(row=lin, column=oc + 5, value=f"=INDEX(${L(col + 4)}${r}:${L(col + 4)}${r + n - 1},{pos})")
    return r, oc, n


rm, ocm, nm_ = ranking(8, 3, "Marcas", marcas, MARCA)
rp, ocp, np_ = ranking(8, 14, "Países", paises, PAIS)
rmod, ocmod, nmod = ranking(8, 25, "Modelos", list(modelo_marca), MODELO, extra=modelo_marca)
nome("marca_lider", f"'Cálculos'!${L(ocm + 1)}${rm}")
nome("pais_lider", f"'Cálculos'!${L(ocp + 1)}${rp}")
nome("modelo_lider", f"'Cálculos'!${L(ocmod + 1)}${rmod}")

# mensal (col A..D, a partir da linha 16)
rmes = bloco(1, 16, "Mensal", ["Mês", "Período", "Ano anterior", "Nº mês"])
for k, m in enumerate(MESES):
    lin = rmes + k
    calc.cell(row=lin, column=1, value=m)
    calc.cell(row=lin, column=4, value=k + 1)
    calc.cell(row=lin, column=2, value=f"=SUMIFS({PRECO},{MES},D{lin},{FA},1)").number_format = MOEDA
    calc.cell(row=lin, column=3, value=f'=IF(sel_ano="Todos",NA(),SUMIFS({PRECO},{MES},D{lin},{FANT},1))').number_format = MOEDA

# anual (col A..C, a partir da linha 32)
rano = bloco(1, 32, "Anual (ignora filtro de ano)", ["Ano", "Faturamento", "Vendas"])
for k, a in enumerate(anos):
    lin = rano + k
    calc.cell(row=lin, column=1, value=str(a))
    calc.cell(row=lin, column=2, value=f"=SUMIFS({PRECO},{ANO},{a},{FSA},1)").number_format = MOEDA
    calc.cell(row=lin, column=3, value=f"=COUNTIFS({ANO},{a},{FSA},1)")
for col in range(1, 22):
    calc.column_dimensions[L(col)].width = 14

# ---------- Dashboard ----------
dash.sheet_view.showGridLines = False
dash.sheet_view.zoomScale = 90
for col in range(1, 21):
    dash.column_dimensions[L(col)].width = 10.5
dash.column_dimensions["A"].width = 2
fundo_banner = PatternFill("solid", fgColor=PRETO)
for row in dash["A1:T3"]:
    for c in row:
        c.fill = fundo_banner
dash["B1"].value = "DASHBOARD DE VENDAS DE VEÍCULOS"
dash["B1"].font = Font(bold=True, size=20, color="FFFFFF")
dash["B2"].value = f"Base: {N} vendas de {anos[0]} a jun/{anos[-1]} · {len(marcas)} marcas · {len(modelo_marca)} modelos · {len(paises)} países"
dash["B2"].font = Font(size=10, color="BBDEFB")
dash.row_dimensions[1].height = 30

# filtros
filtros = [("B5", "C5", "ANO", "Todos", "=lista_anos", "sel_ano"),
           ("E5", "F5", "PAÍS", "Todos", "=lista_paises", "sel_pais"),
           ("H5", "I5", "MARCA", "Todas", "=lista_marcas", "sel_marca")]
for rot, cel, txt, padrao, fonte, nm in filtros:
    dash[rot].value = txt
    dash[rot].font = Font(bold=True, color=AZUL)
    dash[rot].alignment = Alignment(horizontal="right")
    c = dash[cel]
    c.value = padrao
    c.fill = PatternFill("solid", fgColor=AZUL_CLARO)
    c.font = Font(bold=True, size=12, color=AZUL_ESC)
    c.alignment = Alignment(horizontal="center")
    c.border = Border(left=Side(style="medium", color=AZUL), right=Side(style="medium", color=AZUL),
                      top=Side(style="medium", color=AZUL), bottom=Side(style="medium", color=AZUL))
    dv = DataValidation(type="list", formula1=fonte, allow_blank=False, showErrorMessage=True,
                        errorTitle="Filtro inválido", error="Escolha um valor da lista.")
    dash.add_data_validation(dv)
    dv.add(cel)
    nome(nm, f"Dashboard!${cel[0]}${cel[1:]}")
dash["K5"].value = "◄ Escolha nas listas para filtrar todo o painel"
dash["K5"].font = Font(italic=True, color=CINZA_TXT, size=9)

# cartões de KPI
cards = [("B", "FATURAMENTO", "=kpi_fat", MOEDA_K), ("E", "Nº DE VENDAS", "=kpi_qtd", "#,##0"),
         ("H", "TICKET MÉDIO", "=kpi_ticket", MOEDA), ("K", "VS ANO ANTERIOR", '=IF(kpi_var="","—",kpi_var)', '+0.0%;-0.0%;0.0%'),
         ("N", "MARCA LÍDER", "=marca_lider", None), ("Q", "MODELO LÍDER", "=modelo_lider", None)]
for col, titulo, form, fmt in cards:
    i = openpyxl.utils.column_index_from_string(col)
    c1, c2 = col, L(i + 2)
    for row in dash[f"{c1}7:{c2}10"]:
        for c in row:
            c.fill = PatternFill("solid", fgColor=PRETO if col in ("B", "N") else AZUL_ESC if col in ("E", "Q") else AZUL)
    dash.merge_cells(f"{c1}7:{c2}7")
    dash.merge_cells(f"{c1}8:{c2}10")
    t = dash[f"{c1}7"]
    t.value = titulo
    t.font = Font(bold=True, size=9, color="BBDEFB")
    t.alignment = Alignment(horizontal="center", vertical="bottom")
    v = dash[f"{c1}8"]
    v.value = form
    v.font = Font(bold=True, size=18, color="FFFFFF")
    v.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
    if fmt:
        v.number_format = fmt
dash["K11"].value = '=IF(sel_ano="Todos","Selecione um ano para comparar","Comparado com "&(sel_ano-1))'
dash["K11"].font = Font(italic=True, color=CINZA_TXT, size=8)
dash["N11"].value = '="País líder: "&pais_lider'
dash["N11"].font = Font(italic=True, color=CINZA_TXT, size=8)


def estilo(chart, titulo):
    chart.title = titulo
    chart.style = 2
    chart.legend = None
    chart.height = 8.2
    chart.width = 16.5


def pinta(serie, cor, linha=False):
    if linha:
        serie.graphicalProperties.line.solidFill = cor
        serie.graphicalProperties.line.width = 28575
        serie.smooth = False
    else:
        serie.graphicalProperties.solidFill = cor
        serie.graphicalProperties.line.solidFill = cor


# 1. barras por marca (ordenado)
bar = BarChart()
bar.type = "bar"
estilo(bar, "Faturamento por marca")
bar.add_data(Reference(calc, min_col=ocm + 2, min_row=rm - 1, max_row=rm + nm_ - 1), titles_from_data=True)
bar.set_categories(Reference(calc, min_col=ocm + 1, min_row=rm, max_row=rm + nm_ - 1))
bar.x_axis.scaling.orientation = "maxMin"  # maior no topo
bar.y_axis.numFmt = MOEDA_K
bar.y_axis.majorGridlines = None
bar.gapWidth = 60
pinta(bar.series[0], AZUL)
bar.dataLabels = DataLabelList()
bar.dataLabels.showVal = True
bar.dataLabels.numFmt = MOEDA_K
bar.x_axis.delete = False
bar.y_axis.delete = True
dash.add_chart(bar, "B13")

# 2. mensal: período x ano anterior
ln = LineChart()
estilo(ln, "Faturamento mensal: período x ano anterior")
ln.add_data(Reference(calc, min_col=2, max_col=3, min_row=rmes - 1, max_row=rmes + 11), titles_from_data=True)
ln.set_categories(Reference(calc, min_col=1, min_row=rmes, max_row=rmes + 11))
pinta(ln.series[0], AZUL, linha=True)
pinta(ln.series[1], "9E9E9E", linha=True)
ln.series[1].graphicalProperties.line.dashStyle = "dash"
ln.legend = openpyxl.chart.legend.Legend()
ln.legend.position = "b"
ln.y_axis.numFmt = MOEDA_K
ln.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E0E0E0"))
ln.x_axis.delete = False
ln.y_axis.delete = False
dash.add_chart(ln, "K13")

# 3. rosca por país
dn = DoughnutChart()
estilo(dn, "Participação por país")
dn.add_data(Reference(calc, min_col=ocp + 2, min_row=rp - 1, max_row=rp + np_ - 1), titles_from_data=True)
dn.set_categories(Reference(calc, min_col=ocp + 1, min_row=rp, max_row=rp + np_ - 1))
dn.holeSize = 55
for k in range(np_):
    pt = DataPoint(idx=k)
    pt.graphicalProperties.solidFill = PALETA[k % len(PALETA)]
    pt.graphicalProperties.line.solidFill = "FFFFFF"
    dn.series[0].dPt.append(pt)
dn.dataLabels = DataLabelList()
dn.dataLabels.showPercent = True
dn.legend = openpyxl.chart.legend.Legend()
dn.legend.position = "r"
dash.add_chart(dn, "B30")

# 4. evolução anual
col = BarChart()
estilo(col, "Evolução anual do faturamento")
col.add_data(Reference(calc, min_col=2, min_row=rano - 1, max_row=rano + len(anos) - 1), titles_from_data=True)
col.set_categories(Reference(calc, min_col=1, min_row=rano, max_row=rano + len(anos) - 1))
pinta(col.series[0], AZUL_ESC)
col.gapWidth = 50
col.y_axis.numFmt = MOEDA_K
col.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E0E0E0"))
col.dataLabels = DataLabelList()
col.dataLabels.showVal = True
col.dataLabels.numFmt = MOEDA_K
col.x_axis.delete = False
col.y_axis.delete = False
dash.add_chart(col, "K30")

# 5. top 10 modelos (tabela)
dash["B47"].value = "TOP 10 MODELOS"
dash["B47"].font = Font(bold=True, size=12, color=AZUL)
cabs = [("B", "#"), ("C", "Modelo"), ("E", "Marca"), ("G", "Vendas"), ("H", "Faturamento"), ("J", "% do total")]
for c, h in cabs:
    cel = dash[f"{c}48"]
    cel.value = h
    cel.font = Font(bold=True, color="FFFFFF")
    cel.fill = PatternFill("solid", fgColor=PRETO)
for c in "DFIK":
    dash[f"{c}48"].fill = PatternFill("solid", fgColor=PRETO)
dash.merge_cells("C48:D48")
dash.merge_cells("E48:F48")
dash.merge_cells("H48:I48")
dash.merge_cells("J48:K48")
for k in range(10):
    r = 49 + k
    src_r = rmod + k
    vazio = f"'Cálculos'!{L(ocmod + 2)}{src_r}=0"  # esconde posições sem venda no filtro
    dash[f"B{r}"].value = f'=IF({vazio},"",{k + 1})'
    dash[f"C{r}"].value = f"=IF({vazio},\"\",'Cálculos'!{L(ocmod + 1)}{src_r})"
    dash[f"E{r}"].value = f"=IF({vazio},\"\",'Cálculos'!{L(ocmod + 5)}{src_r})"
    dash[f"G{r}"].value = f"=IF({vazio},\"\",'Cálculos'!{L(ocmod + 3)}{src_r})"
    dash[f"H{r}"].value = f"=IF({vazio},\"\",'Cálculos'!{L(ocmod + 2)}{src_r})"
    dash[f"H{r}"].number_format = MOEDA
    dash[f"J{r}"].value = f"=IF({vazio},\"\",'Cálculos'!{L(ocmod + 4)}{src_r})"
    dash[f"J{r}"].number_format = "0.0%"
    for a, b in (("C", "D"), ("E", "F"), ("H", "I"), ("J", "K")):
        dash.merge_cells(f"{a}{r}:{b}{r}")
    for c in "BCDEFGHIJK":
        dash[f"{c}{r}"].border = Border(bottom=fino)
        if k % 2 == 0:
            dash[f"{c}{r}"].fill = PatternFill("solid", fgColor="F5F9FF")
    dash[f"B{r}"].alignment = Alignment(horizontal="center")
    dash[f"G{r}"].alignment = Alignment(horizontal="center")
dash.conditional_formatting.add("H49:H58", DataBarRule(start_type="num", start_value=0, end_type="max", color="64B5F6"))

# notas
dash["M47"].value = "COMO LER ESTE PAINEL"
dash["M47"].font = Font(bold=True, size=12, color=AZUL)
notas = ["• Os filtros de Ano, País e Marca atualizam todos os números e gráficos.",
         "• “Vs ano anterior” e a linha cinza só aparecem com um ano selecionado.",
         "• A evolução anual respeita País e Marca, mas mostra todos os anos.",
         f"• {anos[-1]} tem dados só até junho: compare com cuidado.",
         "• Dados fictícios, usados apenas para estudo."]
for k, t in enumerate(notas):
    dash[f"M{49 + k}"].value = t
    dash[f"M{49 + k}"].font = Font(size=9, color="424242")

dash.conditional_formatting.add("K8", CellIsRule(operator="lessThan", formula=["0"], font=Font(bold=True, size=18, color="FFCDD2")))
dash.conditional_formatting.add("K8", CellIsRule(operator="greaterThan", formula=["0"], font=Font(bold=True, size=18, color="C8E6C9")))

for ws, cor in ((dash, PRETO), (base, AZUL), (calc, "9E9E9E")):
    ws.sheet_properties.tabColor = cor
wb.active = 0
wb.save(OUT)
print("ok", N, "linhas")
